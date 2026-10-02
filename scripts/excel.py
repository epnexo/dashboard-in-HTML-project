"""Excel version of the report: Summary, Charts, Data and Cleaning sheets.

The indicators and tables in Summary are formulas over the Data sheet, so they
recalculate if the data is edited. Needs openpyxl (and pillow for the logo).
"""
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = 'Arial'
TEAL, GREEN, RED, ORANGE = '0D9488', '0CA30C', 'D03B3B', 'EC835A'
INT, PCT = '#,##0', '0.0%'
MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
DATA_COLS = ['transaction_id', 'txn_date', 'customer_id', 'region', 'customer_segment', 'product_type',
             'bundle_name', 'channel', 'data_volume_mb', 'amount', 'status', 'amount_flag']

F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=16, bold=True)
F_SECTION = Font(name=FONT, size=12, bold=True)
F_GREY = Font(name=FONT, size=10, color='4A5568')
F_HEAD = Font(name=FONT, size=10, bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', start_color='162C4C')
LINE = Border(bottom=Side(style='thin', color='E4E8EE'))


def _header(ws, row, titles, col=1):
    for i, t in enumerate(titles):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font, c.fill = F_HEAD, HEAD_FILL
        c.alignment = Alignment(horizontal='left' if i == 0 else 'right', vertical='center', wrap_text=True)


def _table(ws, row, title, titles, rows, formats):
    """Writes title, header and rows. Returns (header row, last row, next free row)."""
    ws.cell(row=row, column=1, value=title).font = F_SECTION
    head = row + 1
    _header(ws, head, titles)
    for i, values in enumerate(rows):
        for j, v in enumerate(values):
            c = ws.cell(row=head + 1 + i, column=1 + j, value=v)
            c.font, c.border = F_BASE, LINE
            if j and formats[j - 1]:
                c.number_format = formats[j - 1]
    last = head + len(rows)
    return head, last, last + 3


def _bars(ws, title, head, last, value_col=2):
    g = BarChart()
    g.type = 'bar'
    g.title, g.legend = title, None
    g.add_data(Reference(ws, min_col=value_col, min_row=head, max_row=last), titles_from_data=True)
    g.set_categories(Reference(ws, min_col=1, min_row=head + 1, max_row=last))
    g.series[0].graphicalProperties.solidFill = TEAL
    g.series[0].graphicalProperties.line.solidFill = TEAL
    g.gapWidth = 60
    g.x_axis.delete = g.y_axis.delete = False
    g.y_axis.number_format = '#,##0'
    g.y_axis.majorGridlines = None
    g.x_axis.scaling.orientation = 'maxMin'  # first category on top
    g.y_axis.crosses = 'max'  # and the value axis at the bottom
    g.title.overlay = False
    g.width, g.height = 16, 8
    return g


def build_excel(out, log, groups, target, logo=None, settings=None):
    settings = settings or {}
    symbol = settings.get('currency_symbol', '$')
    MONEY = f'"{symbol}"#,##0.00'
    n = len(out)
    R = lambda col: f'Data!${col}$2:${col}${n + 1}'  # noqa: E731
    AMOUNT, STATUS, CUSTOMER, MONTH = R('J'), R('K'), R('C'), R('M')

    wb = Workbook()
    summ = wb.active
    summ.title = 'Summary'
    charts, data, cln = wb.create_sheet('Charts'), wb.create_sheet('Data'), wb.create_sheet('Cleaning')

    # ---------- Data ----------
    _header(data, 1, DATA_COLS + ['month'])
    for i, r in enumerate(out, start=2):
        row = [r[c] for c in DATA_COLS]
        row[1] = date.fromisoformat(r['txn_date']) if r['txn_date'] else None
        row[8] = int(r['data_volume_mb']) if r['data_volume_mb'] else None
        row[9] = float(r['amount'])
        row.append(f'=IF(B{i}="","",MONTH(B{i}))')
        data.append(row)
        data.cell(row=i, column=2).number_format = 'yyyy-mm-dd'
        data.cell(row=i, column=10).number_format = '#,##0.00'
    for row in data.iter_rows(min_row=2):
        for c in row:
            c.font = F_BASE
    data.freeze_panes = 'A2'
    data.auto_filter.ref = f'A1:M{n + 1}'
    for i, width in enumerate([14, 12, 13, 15, 18, 14, 20, 12, 16, 10, 12, 18, 7], start=1):
        data.column_dimensions[get_column_letter(i)].width = width

    # ---------- Summary ----------
    # brand band: logo in A1 and the title to its right (or in A1 when there is no logo)
    band = PatternFill('solid', start_color='0F1D37')
    for row in summ['A1:E2']:
        for c in row:
            c.fill = band
    summ.row_dimensions[1].height, summ.row_dimensions[2].height = 30, 27
    col = 'A'
    if logo and logo.exists():
        try:
            from openpyxl.drawing.image import Image
            summ.add_image(Image(str(logo)), 'A1')
            col = 'B'
        except ImportError:  # pillow missing
            pass
    summ[col + '1'] = 'Telecom Transactions Report'
    summ[col + '1'].font = Font(name=FONT, size=16, bold=True, color='FFFFFF')
    summ[col + '1'].alignment = Alignment(vertical='bottom')
    dates = sorted(r['txn_date'] for r in out if r['txn_date'])
    summ[col + '2'] = f'Source: {log["file"]} · Period: {dates[0]} to {dates[-1]} · Currency: {settings.get("currency_name", "")} ({symbol})'
    summ[col + '2'].font = Font(name=FONT, size=10, color='A9B6CC')
    summ[col + '2'].alignment = Alignment(vertical='top')

    # KPIs live in B6:B13; several formulas below reference those cells
    kpis = [
        ('Revenue', f'=SUMIFS({AMOUNT},{STATUS},"Successful")', MONEY, 'Sum of amounts of successful transactions'),
        ('Transactions', f'=COUNTA({R("A")})', INT, 'All statuses'),
        ('Successful transactions', f'=COUNTIF({STATUS},"Successful")', INT, ''),
        ('Success rate', '=IF(B7=0,0,B8/B7)', PCT, 'Successful / transactions'),
        ('Active customers', f'=SUMPRODUCT(1/COUNTIF({CUSTOMER},{CUSTOMER}))', INT, 'Distinct customers with any transaction'),
        ('Average ticket', '=IF(B8=0,0,B6/B8)', MONEY, 'Revenue / successful transactions'),
        ('Revenue per customer', f'=B6/SUMPRODUCT(({STATUS}="Successful")/COUNTIFS({CUSTOMER},{CUSTOMER},{STATUS},{STATUS}))',
         MONEY, 'Revenue / customers with at least one successful purchase'),
        ('Reversed amount', f'=-SUMIFS({AMOUNT},{STATUS},"Reversed")', MONEY, 'Sum of reversed transactions'),
    ]
    _, _, nxt = _table(summ, 4, 'Indicators', ['Indicator', 'Value', 'How it is calculated'],
                       [(k, f, note) for k, f, _, note in kpis], [None, None])
    for i, (_, _, fmt, _) in enumerate(kpis):
        c = summ.cell(row=6 + i, column=2)
        c.number_format, c.font = fmt, F_BOLD
        summ.cell(row=6 + i, column=3).font = F_GREY
        summ.cell(row=6 + i, column=3).alignment = Alignment(horizontal='left')
    summ.cell(row=5, column=3).alignment = Alignment(horizontal='left')

    # the order of each table is fixed by revenue when the file is built
    revenue = defaultdict(lambda: defaultdict(float))
    for r in out:
        for dim in ('product_type', 'region', 'channel', 'customer_segment', 'bundle_name'):
            revenue[dim][r[dim]] += float(r['amount']) if r['status'] == 'Successful' else 0
    order = {dim: sorted(v, key=lambda k: -v[k]) for dim, v in revenue.items()}
    marks = {}

    rows = [(MONTHS[m - 1], f'=SUMIFS({AMOUNT},{STATUS},"Successful",{MONTH},{m})',
             f'=COUNTIFS({STATUS},"Successful",{MONTH},{m})') for m in range(1, 13)]
    head, last, nxt = _table(summ, nxt, 'Monthly revenue', ['Month', 'Revenue', 'Successful transactions'], rows, [MONEY, INT])
    marks['month'] = (head, last)
    summ.cell(row=last + 1, column=1, value=f'The {log["dates_missing"]} transactions without a date are not in this table.').font = F_GREY

    for key, title, label, col_letter in [('product_type', 'Revenue by product type', 'Product', 'F'),
                                          ('region', 'Revenue by region', 'Region', 'D'),
                                          ('channel', 'Revenue by channel', 'Channel', 'H'),
                                          ('customer_segment', 'Revenue by segment', 'Segment', 'E'),
                                          ('bundle_name', 'Bundle performance', 'Bundle', 'G')]:
        first = nxt + 2
        rows = [(v, f'=SUMIFS({AMOUNT},{STATUS},"Successful",{R(col_letter)},A{first + i})',
                 f'=COUNTIFS({STATUS},"Successful",{R(col_letter)},A{first + i})',
                 f'=IF($B$6=0,0,B{first + i}/$B$6)') for i, v in enumerate(order[key])]
        head, last, nxt = _table(summ, nxt, title, [label, 'Revenue', 'Successful transactions', '% of revenue'],
                                 rows, [MONEY, INT, PCT])
        marks[key] = (head, last)

    first = nxt + 2
    chan = R('H')
    rows = [(v, f'=COUNTIFS({chan},A{first + i},{STATUS},"Successful")', f'=COUNTIFS({chan},A{first + i},{STATUS},"Failed")',
             f'=COUNTIFS({chan},A{first + i},{STATUS},"Reversed")',
             f'=IF(SUM(B{first + i}:D{first + i})=0,0,B{first + i}/SUM(B{first + i}:D{first + i}))') for i, v in enumerate(order['channel'])]
    head, last, nxt = _table(summ, nxt, 'Transaction outcome by channel',
                             ['Channel', 'Successful', 'Failed', 'Reversed', 'Success rate'], rows, [INT, INT, INT, PCT])
    marks['status'] = (head, last)
    for letter, width in zip('ABCDE', [30, 16, 46, 16, 16]):
        summ.column_dimensions[letter].width = width
    summ.freeze_panes = 'A3'

    # ---------- Charts ----------
    charts['A1'] = 'Charts'
    charts['A1'].font = F_TITLE
    charts['A2'] = 'Fed by the tables in the Summary sheet.'
    charts['A2'].font = F_GREY
    line = LineChart()
    line.title, line.legend = f'Monthly revenue ({symbol})', None
    head, last = marks['month']
    line.add_data(Reference(summ, min_col=2, min_row=head, max_row=last), titles_from_data=True)
    line.set_categories(Reference(summ, min_col=1, min_row=head + 1, max_row=last))
    line.series[0].graphicalProperties.line.solidFill = TEAL
    line.series[0].graphicalProperties.line.width = 25000
    line.series[0].smooth = False
    line.x_axis.delete = line.y_axis.delete = False
    line.y_axis.number_format = '#,##0'
    line.y_axis.scaling.min = 0
    line.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='E4E8EE'))
    line.title.overlay = False
    line.width, line.height = 33, 8
    charts.add_chart(line, 'A4')

    pos = iter(['A21', 'K21', 'A38', 'K38', 'A55'])
    for key, title in [('product_type', f'Revenue by product type ({symbol})'), ('region', f'Revenue by region ({symbol})'),
                       ('channel', f'Revenue by channel ({symbol})'), ('bundle_name', f'Revenue by bundle ({symbol})')]:
        charts.add_chart(_bars(summ, title, *marks[key]), next(pos))

    head, last = marks['status']
    st = BarChart()
    st.type, st.grouping, st.overlap = 'bar', 'percentStacked', 100
    st.title = 'Transaction outcome by channel'
    st.add_data(Reference(summ, min_col=2, max_col=4, min_row=head, max_row=last), titles_from_data=True)
    st.set_categories(Reference(summ, min_col=1, min_row=head + 1, max_row=last))
    for series, color in zip(st.series, (GREEN, RED, ORANGE)):
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line.solidFill = color
    st.x_axis.delete = st.y_axis.delete = False
    st.x_axis.scaling.orientation = 'maxMin'
    st.y_axis.crosses = 'max'
    st.y_axis.majorGridlines = None
    st.y_axis.number_format = '0%'
    st.title.overlay = False
    st.legend.position = 'b'
    st.legend.overlay = False
    st.gapWidth = 60
    st.width, st.height = 16, 8
    charts.add_chart(st, next(pos))

    # ---------- Cleaning ----------
    cln['A1'] = 'Cleaning applied'
    cln['A1'].font = F_TITLE
    cln['A2'] = (f'{log["rows_in"]} rows in the original file, {log["duplicates_removed"]} duplicates removed, '
                 f'{log["rows_out"]} clean rows. Row-by-row detail in cleaning/findings_detail.csv.')
    cln['A2'].font = F_GREY
    _, _, nxt = _table(cln, 4, 'What was found and what was done',
                       ['Column', 'Issue found', 'Rows', 'Examples of the original value', 'What was done'],
                       [(g['column'], g['issue'], g['n'], ' · '.join(g['examples']), g['action']) for g in groups],
                       [None, INT, None, None])
    _table(cln, nxt, 'Outlier amounts corrected', ['Transaction', 'Bundle', 'Status', 'Original amount', 'Corrected amount'],
           [(a['transaction_id'], a['bundle'], a['status'], a['original_amount'], a['corrected_amount']) for a in log['outliers']],
           [None, None, MONEY, MONEY])
    for row in cln.iter_rows(min_row=6, max_row=5 + len(groups)):
        for c in row:
            c.alignment = Alignment(horizontal='right' if c.column == 3 else 'left', wrap_text=True, vertical='top')
    for letter, width in zip('ABCDE', [16, 46, 16, 50, 46]):
        cln.column_dimensions[letter].width = width

    wb.save(target)
